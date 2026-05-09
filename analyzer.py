import os
import ast
import sqlite3
from database import get_conn, DB_PATH

# ─────────────────────────────────────────
# تحليل Python
# ─────────────────────────────────────────
def analyze_python_file(filepath):
    try:
        with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
            content = f.read()

        tree = ast.parse(content)
        lines = content.splitlines()

        functions   = [n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef)]
        classes     = [n for n in ast.walk(tree) if isinstance(n, ast.ClassDef)]
        imports_std = [n for n in ast.walk(tree) if isinstance(n, ast.Import)]
        imports_frm = [n for n in ast.walk(tree) if isinstance(n, ast.ImportFrom)]

        # استخراج أسماء المكتبات
        libs = set()
        for n in imports_std:
            for alias in n.names:
                libs.add(alias.name.split('.')[0])
        for n in imports_frm:
            if n.module:
                libs.add(n.module.split('.')[0])

        # عدد التعليقات
        comment_lines = sum(1 for l in lines if l.strip().startswith('#'))

        # اكتشاف الأنماط
        patterns = []
        if classes:
            patterns.append('OOP')
        if 'sqlite3' in libs or 'sqlalchemy' in libs:
            patterns.append('قاعدة بيانات')
        if 'tkinter' in libs:
            patterns.append('واجهة رسومية')
        if 'flask' in libs or 'django' in libs or 'fastapi' in libs:
            patterns.append('تطبيق ويب')
        if 'requests' in libs or 'httpx' in libs:
            patterns.append('طلبات HTTP')
        if 'pandas' in libs or 'numpy' in libs:
            patterns.append('معالجة بيانات')

        # حساب متوسط طول الدوال
        func_lengths = []
        for fn in functions:
            length = (fn.end_lineno - fn.lineno + 1) if hasattr(fn, 'end_lineno') else 0
            func_lengths.append(length)
        avg_func_length = round(sum(func_lengths) / len(func_lengths), 1) if func_lengths else 0

        return {
            'lines':           len(lines),
            'comment_lines':   comment_lines,
            'functions':       len(functions),
            'function_names':  [f.name for f in functions],
            'classes':         len(classes),
            'class_names':     [c.name for c in classes],
            'imports':         len(imports_std) + len(imports_frm),
            'libraries':       sorted(libs),
            'patterns':        patterns,
            'avg_func_length': avg_func_length,
            'has_errors':      False,
        }
    except SyntaxError:
        return {'has_errors': True, 'functions': 0, 'function_names': [],
                'classes': 0, 'class_names': [], 'imports': 0,
                'lines': 0, 'comment_lines': 0, 'libraries': [],
                'patterns': [], 'avg_func_length': 0}
    except Exception:
        return None


# ─────────────────────────────────────────
# تحليل PDF
# ─────────────────────────────────────────
def analyze_pdf_file(filepath):
    try:
        from PyPDF2 import PdfReader
        reader   = PdfReader(filepath)
        pages    = len(reader.pages)

        # استخراج النص من أول 3 صفحات
        text = ''
        for page in reader.pages[:3]:
            try:
                text += page.extract_text() or ''
            except Exception:
                pass

        words      = len(text.split()) if text.strip() else 0
        has_images = any(
            '/XObject' in str(page.get('/Resources', ''))
            for page in reader.pages
        )

        return {
            'pages':      pages,
            'words':      words,
            'has_images': has_images,
            'has_text':   words > 50,
        }
    except Exception:
        return None


# ─────────────────────────────────────────
# تحليل Word
# ─────────────────────────────────────────
def analyze_docx_file(filepath):
    try:
        from docx import Document
        doc = Document(filepath)

        paragraphs  = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        word_count  = sum(len(p.split()) for p in paragraphs)
        headings    = [p.text for p in doc.paragraphs
                       if p.style.name.startswith('Heading')]
        table_count = len(doc.tables)

        return {
            'paragraphs':  len(paragraphs),
            'words':       word_count,
            'headings':    headings[:10],
            'tables':      table_count,
        }
    except Exception:
        return None


# ─────────────────────────────────────────
# التحليل الكامل للمشروع
# ─────────────────────────────────────────
def analyze_project(project_id):
    conn = get_conn()
    c    = conn.cursor()

    c.execute("SELECT * FROM projects WHERE id=?", (project_id,))
    project = c.fetchone()

    c.execute("SELECT * FROM files WHERE project_id=?", (project_id,))
    files = c.fetchall()
    conn.close()

    if not project:
        return None

    # ── تجميع الإحصائيات ──
    total_lines = total_functions = total_classes = 0
    total_comments = 0
    all_functions = []
    all_classes   = []
    all_libraries = set()
    all_patterns  = set()
    python_files  = 0
    pdf_pages     = 0
    pdf_words     = 0
    pdf_has_imgs  = False
    docx_paras    = 0
    docx_words    = 0
    docx_tables   = 0
    docx_headings = []
    other_files   = 0
    has_errors    = False

    for f in files:
        filepath = f[3]
        filename = f[2]

        if filename.endswith('.py'):
            result = analyze_python_file(filepath)
            if result:
                python_files    += 1
                total_lines     += result['lines']
                total_functions += result['functions']
                total_classes   += result['classes']
                total_comments  += result['comment_lines']
                all_functions.extend(result['function_names'])
                all_classes.extend(result['class_names'])
                all_libraries.update(result['libraries'])
                all_patterns.update(result['patterns'])
                if result['has_errors']:
                    has_errors = True

        elif filename.endswith('.pdf'):
            result = analyze_pdf_file(filepath)
            if result:
                pdf_pages    += result['pages']
                pdf_words    += result['words']
                pdf_has_imgs  = pdf_has_imgs or result['has_images']

        elif filename.endswith('.docx'):
            result = analyze_docx_file(filepath)
            if result:
                docx_paras   += result['paragraphs']
                docx_words   += result['words']
                docx_tables  += result['tables']
                docx_headings.extend(result['headings'])

        else:
            other_files += 1

    # ── نسبة الجاهزية ──
    readiness = 100
    if python_files == 0:           readiness -= 30
    if total_functions < 5:         readiness -= 15
    if total_classes < 1:           readiness -= 10
    if pdf_pages == 0 and docx_paras == 0: readiness -= 20
    if has_errors:                  readiness -= 30
    if total_comments == 0:         readiness -= 5
    readiness = max(readiness, 10)

    # ── سؤال مقترح ذكي ──
    suggested_question = _generate_question(
        all_functions, all_classes, all_patterns, all_libraries
    )

    return {
        'project':          project,
        'python_files':     python_files,
        'total_lines':      total_lines,
        'total_functions':  total_functions,
        'total_classes':    total_classes,
        'total_comments':   total_comments,
        'all_functions':    all_functions[:8],
        'all_classes':      all_classes[:8],
        'libraries':        sorted(all_libraries),
        'patterns':         sorted(all_patterns),
        'pdf_pages':        pdf_pages,
        'pdf_words':        pdf_words,
        'pdf_has_images':   pdf_has_imgs,
        'docx_paragraphs':  docx_paras,
        'docx_words':       docx_words,
        'docx_tables':      docx_tables,
        'docx_headings':    docx_headings[:5],
        'other_files':      other_files,
        'readiness':        readiness,
        'has_errors':       has_errors,
        'suggested_question': suggested_question,
        'files':            files,
    }


def _generate_question(functions, classes, patterns, libraries):
    """توليد سؤال مناقشة مخصص بناءً على محتوى المشروع"""
    if 'قاعدة بيانات' in patterns and classes:
        return f"كيف ربطت كلاس '{classes[0]}' بقاعدة البيانات؟ واشرح عمليات CRUD في مشروعك."
    if 'واجهة رسومية' in patterns and functions:
        fn = max(functions, key=len)
        return f"اشرح كيف تتعامل دالة '{fn}' مع أحداث الواجهة الرسومية."
    if 'تطبيق ويب' in patterns:
        return "ما هي المسارات (routes) الرئيسية في تطبيقك وكيف تتعامل مع طلبات المستخدم؟"
    if 'OOP' in patterns and classes:
        return f"اشرح مبدأ التوارث أو التغليف في كلاس '{classes[0]}' وكيف طبقته."
    if functions:
        fn = max(functions, key=len)
        return f"كيف تعمل دالة '{fn}' وما دورها الأساسي في النظام؟"
    return "ما هي أبرز التحديات التقنية التي واجهتها وكيف تغلبت عليها؟"
